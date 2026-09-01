#!/usr/bin/env python3
"""Build a supporting-material DOCX from a reviewed JSON content map.

The JSON map is semantic input prepared from the final canonical English. This
script controls presentation only; it never translates, rewrites, or invents
missing lesson metadata.
"""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROFILE_ID = "MPI-SM-TPL 1.0.0"
FIXED_HEADINGS = {
    "preliminary_reflection": "Preliminary Reflection",
    "study_requirements": "Study Requirements",
    "reflection_questions": "Reflection Questions",
    "contemplation_questions": "Contemplation Questions",
    "learning_application": "Learning Application",
    "sharing_appreciation": "Sharing Appreciation",
    "meditation_guidance": "Meditation Guidance",
    "self_reflection": "Self-Reflection",
}
FIXED_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)


def save_deterministic_docx(document, output_path):
    """Save stable OOXML bytes by normalizing package order and metadata."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    package = BytesIO()
    document.save(package)
    package.seek(0)
    with zipfile.ZipFile(package, "r") as source, zipfile.ZipFile(
        output_path, "w", compression=zipfile.ZIP_STORED
    ) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIMESTAMP)
            info.compress_type = zipfile.ZIP_STORED
            info.create_system = 0
            target.writestr(info, source.read(name))


def set_run_font(run, name="Times New Roman", size=12, *, bold=None, italic=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{key}"), name)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:bidi"), "en-US")


def set_keep(paragraph, *, next_paragraph=False, lines=False):
    ppr = paragraph._p.get_or_add_pPr()
    for tag, enabled in (("w:keepNext", next_paragraph), ("w:keepLines", lines)):
        if enabled and ppr.find(qn(tag)) is None:
            ppr.append(OxmlElement(tag))


def set_body_spacing(paragraph, *, before=0, after=5, line=1.08):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line
    fmt.widow_control = True


def add_numbering_definition(doc, levels):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multilevel = OxmlElement("w:multiLevelType")
    multilevel.set(qn("w:val"), "multilevel")
    abstract.append(multilevel)

    for ilvl, spec in enumerate(levels):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        for tag, value in (
            ("w:start", spec.get("start", 1)),
            ("w:numFmt", spec["format"]),
            ("w:lvlText", spec["text"]),
            ("w:suff", "space"),
        ):
            node = OxmlElement(tag)
            node.set(qn("w:val"), str(value))
            lvl.append(node)
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(spec["left"]))
        tabs.append(tab)
        ppr.append(tabs)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), str(spec["left"]))
        indent.set(qn("w:hanging"), str(spec["hanging"]))
        ppr.append(indent)
        lvl.append(ppr)
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), str(spec.get("font", "Times New Roman")))
        fonts.set(qn("w:hAnsi"), str(spec.get("font", "Times New Roman")))
        rpr.append(fonts)
        lvl.append(rpr)
        abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id, level=0):
    ppr = paragraph._p.get_or_add_pPr()
    old = ppr.find(qn("w:numPr"))
    if old is not None:
        ppr.remove(old)
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    numpr.append(ilvl)
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(numid)
    ppr.insert(0, numpr)


def add_page_field(paragraph, code):
    run = paragraph.add_run()
    set_run_font(run, "Arial", 9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {code} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def configure_document(doc):
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(14)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    heading = doc.styles["Heading 1"]
    heading.font.name = "Times New Roman"
    heading.font.size = Pt(14)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("Page ")
    set_run_font(run, "Arial", 9)
    add_page_field(paragraph, "PAGE")
    run = paragraph.add_run(" of ")
    set_run_font(run, "Arial", 9)
    add_page_field(paragraph, "NUMPAGES")


def create_base_template(include_skeleton=True):
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "MPI Supporting Material Template"
    doc.core_properties.subject = PROFILE_ID
    doc.core_properties.author = "Mindful Peace International"
    doc.core_properties.keywords = "supporting material, Word template, MPI-SM-TPL"
    if include_skeleton:
        for heading in FIXED_HEADINGS.values():
            add_heading(doc, heading)
            paragraph = doc.add_paragraph(style="Normal")
            set_body_spacing(paragraph, after=3)
    return doc


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def parse_inline(text):
    text = str(text).replace("---", "—")
    segments = []
    cursor = 0
    for match in re.finditer(r"\*([^*]+)\*", text):
        if match.start() > cursor:
            segments.append((text[cursor : match.start()], False))
        segments.append((match.group(1), True))
        cursor = match.end()
    if cursor < len(text):
        segments.append((text[cursor:], False))
    return segments or [(text, False)]


def normalize_meditation_title(text):
    visible = "".join(segment for segment, _ in parse_inline(text)).strip()
    if visible.startswith("[") and visible.endswith("]"):
        visible = visible[1:-1].strip()
    return visible


def add_rich_text(paragraph, text, *, bold=False, honor_italics=True):
    for segment, italic in parse_inline(text):
        run = paragraph.add_run(segment)
        set_run_font(run, bold=bold, italic=bool(italic and honor_italics))


def add_heading(doc, text):
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.0
    set_keep(paragraph, next_paragraph=True, lines=True)
    run = paragraph.add_run(f"[{text}]")
    set_run_font(run, size=14, bold=True)
    return paragraph


def add_body(doc, text, *, italic=False, bold=False, left=0.17, after=5):
    paragraph = doc.add_paragraph(style="Normal")
    set_body_spacing(paragraph, after=after)
    paragraph.paragraph_format.left_indent = Inches(left)
    if italic or bold:
        run = paragraph.add_run(str(text).replace("---", "—"))
        set_run_font(run, bold=bold, italic=italic)
    else:
        add_rich_text(paragraph, text)
    return paragraph


def add_list_item(
    doc,
    text,
    num_id,
    *,
    level=0,
    bold=False,
    keep_next=False,
    honor_italics=True,
):
    paragraph = doc.add_paragraph(style="Normal")
    apply_numbering(paragraph, num_id, level)
    set_body_spacing(paragraph, after=4)
    set_keep(paragraph, next_paragraph=keep_next, lines=True)
    add_rich_text(paragraph, text, bold=bold, honor_italics=honor_italics)
    return paragraph


def heading_for(section):
    role = section["role"]
    if role == "custom":
        heading = str(section.get("heading", "")).strip()
        if not heading:
            raise ValueError("custom section requires a non-empty heading")
        return heading
    if role not in FIXED_HEADINGS:
        raise ValueError(f"unsupported section role: {role}")
    return FIXED_HEADINGS[role]


def require_list(value, label):
    if not isinstance(value, list):
        raise ValueError(f"{label} must be a list")
    return value


def validate_content_map(data):
    if not isinstance(data, dict):
        raise ValueError("content JSON root must be an object")
    if not str(data.get("title", "")).strip():
        raise ValueError("title is required")
    sections = require_list(data.get("sections"), "sections")
    if not sections:
        raise ValueError("sections must not be empty")
    for index, section in enumerate(sections):
        if not isinstance(section, dict) or not section.get("role"):
            raise ValueError(f"sections[{index}] requires a role")
        role = section["role"]
        heading_for(section)
        if role in {"preliminary_reflection", "sharing_appreciation"}:
            content = require_list(section.get("paragraphs", []), f"sections[{index}].paragraphs")
        elif role in {
            "study_requirements",
            "reflection_questions",
            "contemplation_questions",
            "learning_application",
            "self_reflection",
        }:
            content = require_list(section.get("items", []), f"sections[{index}].items")
        elif role == "meditation_guidance":
            content = require_list(section.get("groups", []), f"sections[{index}].groups")
        else:
            paragraphs = require_list(
                section.get("paragraphs", []), f"sections[{index}].paragraphs"
            )
            items = require_list(section.get("items", []), f"sections[{index}].items")
            content = paragraphs + items
        if not content:
            raise ValueError(f"sections[{index}] is empty; omit empty sections from the map")


def build_document(data, template_path):
    validate_content_map(data)
    doc = Document(template_path)
    clear_body(doc)
    configure_document(doc)
    doc.core_properties.title = str(data["title"])
    doc.core_properties.subject = f"Supporting material; {PROFILE_ID}"
    doc.core_properties.author = "MPI translation workflow — AI-reviewed draft"

    lesson_number = data.get("lesson_number")
    has_lesson = lesson_number is not None and bool(str(lesson_number).strip())
    if has_lesson:
        lesson = doc.add_paragraph(style="Title")
        lesson.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lesson.paragraph_format.space_before = Pt(4)
        lesson.paragraph_format.space_after = Pt(2)
        set_keep(lesson, next_paragraph=True, lines=True)
        run = lesson.add_run(f"Lesson {str(lesson_number).strip()}")
        set_run_font(run, size=14, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(4 if not has_lesson else 0)
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.line_spacing = 1.0
    set_keep(title, next_paragraph=True, lines=True)
    run = title.add_run(str(data["title"]).strip())
    set_run_font(run, size=14, bold=True)

    decimal_spec = [{"format": "decimal", "text": "%1.", "left": 360, "hanging": 300}]
    letter_spec = [{"format": "lowerLetter", "text": "%1)", "left": 540, "hanging": 360}]
    bullet_spec = [{"format": "bullet", "text": "•", "left": 540, "hanging": 300, "font": "Arial"}]
    reflection_spec = [
        {"format": "decimal", "text": "%1.", "left": 360, "hanging": 300},
        {"format": "lowerLetter", "text": "(%2)", "left": 720, "hanging": 360},
    ]

    for section in data["sections"]:
        role = section["role"]
        add_heading(doc, heading_for(section))

        if role in {"preliminary_reflection", "sharing_appreciation"}:
            for text in require_list(section.get("paragraphs", []), f"{role}.paragraphs"):
                paragraph = add_body(doc, text, after=7)
                set_keep(paragraph, lines=True)
        elif role == "study_requirements":
            num_id = add_numbering_definition(doc, bullet_spec)
            for text in require_list(section.get("items", []), "study_requirements.items"):
                add_list_item(doc, text, num_id)
        elif role == "reflection_questions":
            num_id = add_numbering_definition(doc, reflection_spec)
            for item in require_list(section.get("items", []), "reflection_questions.items"):
                add_list_item(
                    doc,
                    item["text"],
                    num_id,
                    level=0,
                    bold=True,
                    keep_next=True,
                    honor_italics=False,
                )
                for subquestion in require_list(item.get("subquestions", []), "reflection subquestions"):
                    add_list_item(doc, subquestion, num_id, level=1)
        elif role in {"contemplation_questions", "learning_application", "self_reflection"}:
            num_id = add_numbering_definition(doc, decimal_spec)
            for text in require_list(section.get("items", []), f"{role}.items"):
                add_list_item(doc, text, num_id)
        elif role == "meditation_guidance":
            group_num = add_numbering_definition(doc, decimal_spec)
            body_bullets = add_numbering_definition(doc, bullet_spec)
            for group in require_list(section.get("groups", []), "meditation_guidance.groups"):
                add_list_item(
                    doc,
                    normalize_meditation_title(group["title"]),
                    group_num,
                    bold=True,
                    keep_next=True,
                    honor_italics=False,
                )
                letter_num = add_numbering_definition(doc, letter_spec)
                for key, default_label in (
                    ("contemplative", "Contemplative Meditation"),
                    ("abiding", "Abiding Meditation"),
                ):
                    block = group.get(key)
                    if not block:
                        continue
                    add_list_item(
                        doc,
                        block.get("label", default_label),
                        letter_num,
                        keep_next=True,
                        honor_italics=False,
                    )
                    for text in require_list(block.get("paragraphs", []), f"{key}.paragraphs"):
                        add_list_item(doc, text, body_bullets)
        elif role == "custom":
            for text in require_list(section.get("paragraphs", []), "custom.paragraphs"):
                add_body(doc, text)
            items = require_list(section.get("items", []), "custom.items")
            if items:
                list_kind = section.get("list_kind", "decimal")
                if list_kind == "decimal":
                    num_id = add_numbering_definition(doc, decimal_spec)
                elif list_kind == "bullet":
                    num_id = add_numbering_definition(doc, bullet_spec)
                else:
                    raise ValueError("custom.list_kind must be decimal or bullet")
                for text in items:
                    add_list_item(doc, text, num_id)

    return doc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("content_json", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument(
        "--template",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "assets" / "supporting-material-template-v1.docx",
    )
    args = parser.parse_args()
    data = json.loads(args.content_json.read_text(encoding="utf-8"))
    document = build_document(data, args.template)
    save_deterministic_docx(document, args.output_docx)
    print(f"created {args.output_docx}")
    print(f"profile {PROFILE_ID}")


if __name__ == "__main__":
    main()
