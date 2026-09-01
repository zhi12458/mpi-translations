#!/usr/bin/env python3
"""Validate an MPI-SM-TPL DOCX against its reviewed JSON content map."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from build_supporting_material_docx import heading_for, normalize_meditation_title, parse_inline


FORBIDDEN_PLACEHOLDERS = (
    "Lesson xx",
    "Lesson xxx",
    "Title of the lesson",
    "Key in text here",
    "Question topic",
    "Breakdown 1",
    "Breakdown 2",
)
HAN = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")


def visible_text(value):
    return "".join(text for text, _ in parse_inline(value))


def expected_content(data):
    texts = [str(data["title"]).strip()]
    lesson_number = data.get("lesson_number")
    if lesson_number is not None and str(lesson_number).strip():
        texts.append(f"Lesson {str(lesson_number).strip()}")
    numbered = 0
    headings = []
    for section in data["sections"]:
        role = section["role"]
        headings.append(f"[{heading_for(section)}]")
        if role in {"preliminary_reflection", "sharing_appreciation"}:
            texts.extend(visible_text(text) for text in section.get("paragraphs", []))
        elif role == "study_requirements":
            items = section.get("items", [])
            numbered += len(items)
            texts.extend(visible_text(text) for text in items)
        elif role == "reflection_questions":
            for item in section.get("items", []):
                numbered += 1
                texts.append(visible_text(item["text"]))
                subs = item.get("subquestions", [])
                numbered += len(subs)
                texts.extend(visible_text(text) for text in subs)
        elif role in {"contemplation_questions", "learning_application", "self_reflection"}:
            items = section.get("items", [])
            numbered += len(items)
            texts.extend(visible_text(text) for text in items)
        elif role == "meditation_guidance":
            for group in section.get("groups", []):
                numbered += 1
                texts.append(normalize_meditation_title(group["title"]))
                for key, label in (
                    ("contemplative", "Contemplative Meditation"),
                    ("abiding", "Abiding Meditation"),
                ):
                    block = group.get(key)
                    if not block:
                        continue
                    numbered += 1
                    texts.append(visible_text(block.get("label", label)))
                    paragraphs = block.get("paragraphs", [])
                    numbered += len(paragraphs)
                    texts.extend(visible_text(text) for text in paragraphs)
        elif role == "custom":
            texts.extend(visible_text(text) for text in section.get("paragraphs", []))
            items = section.get("items", [])
            numbered += len(items)
            texts.extend(visible_text(text) for text in items)
    return texts, headings, numbered


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("content_json", type=Path)
    parser.add_argument("--allow-han", action="store_true")
    args = parser.parse_args()

    data = json.loads(args.content_json.read_text(encoding="utf-8"))
    expected, expected_headings, expected_numbered = expected_content(data)
    doc = Document(args.docx)
    visible = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    joined = "\n".join(visible)

    actual_headings = [
        paragraph.text for paragraph in doc.paragraphs if paragraph.style.name == "Heading 1"
    ]
    if actual_headings != expected_headings:
        raise SystemExit(
            f"FAIL heading sequence\nexpected={expected_headings!r}\nactual={actual_headings!r}"
        )

    missing = [text for text in expected if text and text not in visible]
    if missing:
        raise SystemExit(f"FAIL missing content: {missing[:5]!r}")
    leaked = [text for text in FORBIDDEN_PLACEHOLDERS if text in joined]
    if leaked:
        raise SystemExit(f"FAIL template placeholder leakage: {leaked!r}")
    if not args.allow_han and HAN.search(joined):
        raise SystemExit("FAIL visible Han characters found")

    numbered = []
    for paragraph in doc.paragraphs:
        ppr = paragraph._p.pPr
        if ppr is not None and ppr.find(qn("w:numPr")) is not None:
            numbered.append(paragraph.text)
    if len(numbered) != expected_numbered:
        raise SystemExit(
            f"FAIL numbered paragraph count: expected {expected_numbered}, got {len(numbered)}"
        )

    section = doc.sections[0]
    geometry = (
        round(section.page_width.inches, 2),
        round(section.page_height.inches, 2),
        round(section.left_margin.inches, 2),
        round(section.right_margin.inches, 2),
        round(section.top_margin.inches, 2),
        round(section.bottom_margin.inches, 2),
    )
    expected_geometry = (8.5, 11.0, 0.85, 0.85, 0.75, 0.72)
    if geometry != expected_geometry:
        raise SystemExit(f"FAIL geometry: expected {expected_geometry}, got {geometry}")

    with zipfile.ZipFile(args.docx) as package:
        footer = package.read("word/footer1.xml").decode("utf-8")
        if " PAGE " not in footer or " NUMPAGES " not in footer:
            raise SystemExit("FAIL dynamic PAGE/NUMPAGES fields missing")
        bad = package.testzip()
        if bad:
            raise SystemExit(f"FAIL corrupt package member: {bad}")

    print("PASS MPI-SM-TPL document validation")
    print(f"headings={len(expected_headings)}")
    print(f"numbered_paragraphs={expected_numbered}")
    print(f"content_units={len(expected)}")
    print("placeholder_leakage=0")
    print("page_fields=PAGE+NUMPAGES")
    print("geometry=Letter 8.5x11; margins 0.85/0.85/0.75/0.72")


if __name__ == "__main__":
    main()
