#!/usr/bin/env python3
"""Build the human-facing M2/M3 comparison document."""

from __future__ import annotations

import argparse
import difflib
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


DOC_SKILL = Path("/Users/jingzhi/.codex/plugins/cache/openai-primary-runtime/documents/26.819.11345/skills/documents")
sys.path.insert(0, str(DOC_SKILL / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "667085"
SOURCE_FILL = "EEF5FB"
M2_FILL = "FFF4F2"
M3_FILL = "F2F8F3"
M2_DIFF = "B42318"
M3_DIFF = "067647"
M2_HILITE = "FEE4E2"
M3_HILITE = "D1FADF"
TABLE_FILL = "E8EEF5"
GRID = "D0D5DD"


def set_font(run, *, name="Calibri", east_asia="Hiragino Sans GB W3", size=None,
             color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_table_borders(table, color=GRID, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_keep(paragraph, *, next_=False, lines=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if next_ and p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))
    if lines and p_pr.find(qn("w:keepLines")) is None:
        p_pr.append(OxmlElement("w:keepLines"))


def add_field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph._p.extend([begin, instr, separate, text, end])


def clean_djot(text):
    text = re.sub(r"\{#[^{}]*\}", "", text)
    text = re.sub(r"\s*\[\d+\]\(#[^)]*\)", "", text)
    text = re.sub(r"\[([^\[\]]+)\]\(#[^)]*\)", r"\1", text)
    text = re.sub(r"^\s*#{1,6}\s*", "", text)
    text = re.sub(r"^\s*[-*+]\s+", "", text)
    text = text.replace("*", "")
    text = text.replace("------", "—")
    return text.strip()


TOKEN_RE = re.compile(r"\s+|[\w]+(?:['’][\w]+)?|[^\w\s]", re.UNICODE)


def diff_tokens(a, b):
    at = TOKEN_RE.findall(a)
    bt = TOKEN_RE.findall(b)
    matcher = difflib.SequenceMatcher(a=at, b=bt, autojunk=False)
    a_runs, b_runs = [], []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        a_runs.append(("".join(at[i1:i2]), tag != "equal"))
        b_runs.append(("".join(bt[j1:j2]), tag != "equal"))
    return a_runs, b_runs


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB W3")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB W3")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_running_furniture(section):
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("MPI 翻译人工对比稿  |  Great Path, Great Business")
    set_font(r, size=8.5, color=MUTED)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    r = fp.add_run("第 ")
    set_font(r, size=8.5, color=MUTED)
    add_field(fp, "PAGE")
    r = fp.add_run(" 页")
    set_font(r, size=8.5, color=MUTED)


def add_label_paragraph(doc, label, text, fill, label_color, runs=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    set_paragraph_shading(p, fill)
    lr = p.add_run(f"{label}  ")
    set_font(lr, size=9.5, color=label_color, bold=True)
    if runs is None:
        tr = p.add_run(text)
        set_font(tr, size=10.5, color=INK)
    else:
        for chunk, changed in runs:
            if not chunk:
                continue
            tr = p.add_run(chunk)
            set_font(tr, size=10.5, color=label_color if changed else INK,
                     bold=changed)
            if changed:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), M2_HILITE if label.startswith("M2") else M3_HILITE)
                tr._element.get_or_add_rPr().append(shd)
    set_keep(p)
    return p


def add_score_table(doc):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["独立裁判", "M2 稳健版", "M3 风格版", "裁判结论"]
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_fill(cell, TABLE_FILL)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_font(r, size=9.5, color=DARK_BLUE, bold=True)
    rows = [
        ("DeepSeek V4 Pro Max", "92", "93", "M3 +1"),
        ("GPT-5.6 Sol High", "95.6", "96.9", "M3 +1.3"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(value)
            set_font(r, size=9.5, color=INK, bold=idx == 3)
    apply_table_geometry(
        table,
        [2200, 1700, 1700, 3760],
        table_width_dxa=9360,
        indent_dxa=120,
        cell_margins_dxa={"top": 100, "bottom": 100, "start": 120, "end": 120},
    )
    set_table_borders(table)
    return table


def build(source_path, m2_path, m3_path, output_path):
    source = source_path.read_text(encoding="utf-8").splitlines()
    m2 = m2_path.read_text(encoding="utf-8").splitlines()
    m3 = m3_path.read_text(encoding="utf-8").splitlines()
    if not (len(source) == len(m2) == len(m3) == 378):
        raise ValueError("expected three aligned 378-line files")
    changed = [i + 1 for i, (a, b) in enumerate(zip(m2, m3)) if a != b]
    expected = [14, 75, 79, 101, 132, 156, 164, 170, 189, 237, 257, 267,
                271, 273, 285, 287, 290, 304, 320, 344, 354]
    if changed != expected:
        raise ValueError(f"unexpected changed-line set: {changed}")

    doc = Document()
    configure_styles(doc)
    add_running_furniture(doc.sections[0])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("翻译组人工审阅材料")
    set_font(r, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("《大道大商——佛教与企业文化》")
    set_font(r, size=24, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("M2 稳健版 vs M3 限定风格优化版 · 中英彩色差异稿")
    set_font(r, size=13, color=MUTED)

    metadata = [
        ("固定题名", "Great Path, Great Business"),
        ("固定人名", "济群法师 = Master Jiqun"),
        ("对比范围", "仅列出 21 个不同的源文行；其余 357 行完全相同"),
        ("阅读顺序", "每组先中文，后 M2（上）与 M3（下）；彩色底纹标出英文差异"),
        ("交付状态", "AI draft；最终定稿仍需具名佛法审校人确认"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(f"{label}：")
        set_font(r, size=10.5, color=INK, bold=True)
        r = p.add_run(value)
        set_font(r, size=10.5, color=INK)

    doc.add_heading("盲评结果与使用建议", level=1)
    add_score_table(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("建议：")
    set_font(r, size=10.5, color=DARK_BLUE, bold=True)
    r = p.add_run(
        "两位裁判均选择 M3；分数对应 L344 最终微调前的盲评稿。"
        "L344 已另经 V4 Pro 定向复核并清零 blocker，因此 M3 作为 M 1.0 验证基线。"
    )
    set_font(r, size=10.5, color=INK)

    doc.add_heading("颜色说明", level=2)
    add_label_paragraph(doc, "中文", "对应的中文源文", SOURCE_FILL, DARK_BLUE)
    add_label_paragraph(doc, "M2 稳健版", "红色底纹为 M2 中被 M3 改写的文字", M2_FILL, M2_DIFF)
    add_label_paragraph(doc, "M3 风格版", "绿色底纹为 M3 新写或改写的文字", M3_FILL, M3_DIFF)

    doc.add_page_break()
    doc.add_heading("21 处差异逐项对照", level=1)

    for count, line_no in enumerate(changed, start=1):
        heading = doc.add_heading(f"{count:02d}. L{line_no}", level=2)
        set_keep(heading, next_=True)
        src = clean_djot(source[line_no - 1])
        a = clean_djot(m2[line_no - 1])
        b = clean_djot(m3[line_no - 1])
        a_runs, b_runs = diff_tokens(a, b)
        add_label_paragraph(doc, "中文", src, SOURCE_FILL, DARK_BLUE)
        add_label_paragraph(doc, "M2 稳健版", a, M2_FILL, M2_DIFF, a_runs)
        add_label_paragraph(doc, "M3 风格版", b, M3_FILL, M3_DIFF, b_runs)

    doc.core_properties.title = "Great Path, Great Business — M2 vs M3 Human Comparison"
    doc.core_properties.subject = "Bilingual translation comparison"
    doc.core_properties.author = "Mindful Peace International translation workflow"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--m2", type=Path, required=True)
    parser.add_argument("--m3", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.m2, args.m3, args.output)


if __name__ == "__main__":
    main()
